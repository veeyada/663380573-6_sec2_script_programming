# document-automator/src/word_processor.py
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os
from .utils import setup_logging

logger = setup_logging(__name__)

ALIGNMENTS = {
    'left': WD_ALIGN_PARAGRAPH.LEFT,
    'center': WD_ALIGN_PARAGRAPH.CENTER,
    'right': WD_ALIGN_PARAGRAPH.RIGHT,
    'justify': WD_ALIGN_PARAGRAPH.JUSTIFY,
}


class WordProcessor:
    """
    Handles operations related to Word documents using python-docx.
    """

    def __init__(self):
        logger.info("WordProcessor initialized.")

    def create_document(self):
        """Creates a new blank Word document."""
        document = docx.Document()
        logger.info("New Word document created.")
        return document

    def add_heading(self, document, text, level=1):
        """Adds a heading to the document."""
        document.add_heading(text, level=level)
        logger.info(f"Added heading (Level {level}): '{text}'")

    def add_paragraph(self, document, text, style=None, bold=False, italic=False, font_size=None, align='left'):
        """Adds a paragraph with optional formatting."""
        paragraph = document.add_paragraph(text, style=style)

        # Apply run-level formatting if specified
        if paragraph.runs:  # Ensure there's at least one run
            run = paragraph.runs[0]  # Apply to the first run
            if bold:
                run.bold = True
            if italic:
                run.italic = True
            if font_size:
                run.font.size = Pt(font_size)

        # Apply paragraph alignment
        paragraph.alignment = ALIGNMENTS.get(align, WD_ALIGN_PARAGRAPH.LEFT)

        logger.info(f"Added paragraph: '{text[:50]}...'")
        return paragraph

    def add_extracted_text(self, document, text, page_separator="--- Page End ---"):
        """
        Inserts extracted PDF text line by line so the layout stays readable,
        with a sub-heading per page instead of one large paragraph.
        """
        pages = [p for p in text.split(page_separator) if p.strip()]
        for page_num, page_text in enumerate(pages, start=1):
            document.add_heading(f"Page {page_num}", level=3)
            for line in page_text.splitlines():
                line = line.strip()
                if line:
                    document.add_paragraph(line)
        logger.info(f"Inserted extracted text from {len(pages)} page(s).")
        return len(pages)

    def add_table(self, document, data, include_header=True, header_bold=True, col_widths=None):
        """
        Adds a table to the document.
        `data` should be a list of lists (rows).
        `col_widths` is a list of floats (in inches) for column widths.
        """
        if not data:
            logger.warning("No data provided to add table.")
            return None

        rows = len(data)
        cols = len(data[0])

        if rows == 0 or cols == 0:
            logger.warning("Table data is empty or malformed. Skipping table creation.")
            return None

        table = document.add_table(rows=rows, cols=cols)
        table.style = 'Table Grid'  # Apply a basic grid style

        # Populate table
        for r_idx, row_data in enumerate(data):
            cells = table.rows[r_idx].cells
            for c_idx, cell_data in enumerate(row_data):
                cells[c_idx].text = str(cell_data)
                # If header is included in data and we apply special header formatting
                if include_header and header_bold and r_idx == 0:
                    for run in cells[c_idx].paragraphs[0].runs:
                        run.bold = True

        # Set column widths if provided (Word reads the width from each cell)
        if col_widths and len(col_widths) == cols:
            for c_idx, width in enumerate(col_widths):
                for cell in table.columns[c_idx].cells:
                    cell.width = Inches(width)

        logger.info(f"Added table with {rows} rows and {cols} columns.")
        return table

    def add_picture(self, document, image_path, width_inches=None):
        """Adds an image to the document with optional width scaling."""
        if not os.path.exists(image_path):
            logger.error(f"Image file not found: {image_path}")
            return False

        try:
            if width_inches:
                document.add_picture(image_path, width=Inches(width_inches))
            else:
                document.add_picture(image_path)
            logger.info(f"Added picture: {os.path.basename(image_path)}")
            return True
        except Exception as e:
            logger.error(f"Error adding picture '{image_path}': {e}")
            return False

    def add_page_break(self, document):
        """Inserts a page break."""
        document.add_page_break()
        logger.info("Added page break.")

    def read_document(self, docx_path):
        """Reads a .docx and returns (paragraph texts, tables as lists of rows)."""
        if not os.path.exists(docx_path):
            logger.error(f"Word document not found: {docx_path}")
            return None, None
        document = docx.Document(docx_path)
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        tables = [[[cell.text for cell in row.cells] for row in t.rows] for t in document.tables]
        logger.info(f"Read '{docx_path}': {len(paragraphs)} paragraphs, {len(tables)} tables.")
        return paragraphs, tables

    def save_document(self, document, output_path):
        """Saves the Word document to a specified path."""
        try:
            document.save(output_path)
            logger.info(f"Word document saved to '{output_path}'.")
            return True
        except Exception as e:
            logger.error(f"Error saving Word document to '{output_path}': {e}")
            return False
