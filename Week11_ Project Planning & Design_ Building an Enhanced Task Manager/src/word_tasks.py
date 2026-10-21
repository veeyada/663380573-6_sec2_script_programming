# agentic-document-automator/src/word_tasks.py
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os
from utils import setup_logging  # ✅ แก้จาก .utils เป็น utils

logger = setup_logging(__name__)

class WordTasks:
    """Handles advanced Word document operations for the Document Agent."""
    
    def __init__(self):
        logger.info("WordTasks initialized.")

    def _replace_placeholder_in_paragraph(self, paragraph, placeholder, value):
        """Replaces a single placeholder in a paragraph's runs."""
        if placeholder in paragraph.text:
            replaced = False
            for item in paragraph.runs:
                if placeholder in item.text:
                    item.text = item.text.replace(placeholder, str(value))
                    replaced = True
            if not replaced and paragraph.runs:
                # Placeholder is split across several runs (Word does this when text is edited):
                # put the full replaced text in the first run and clear the others.
                full_text = paragraph.text.replace(placeholder, str(value))
                paragraph.runs[0].text = full_text
                for run in paragraph.runs[1:]:
                    run.text = ""
            logger.debug(f"Replaced '{placeholder}' with '{value}' in paragraph.")
            return True  # Placeholder found and replaced
        return False  # Placeholder not found in this paragraph

    @staticmethod
    def _build_table(document, table_data, cols):
        """Create a grid table (header row in bold) at the end of the document and return it."""
        table = document.add_table(rows=0, cols=cols)
        try:
            table.style = 'Table Grid'
        except Exception:
            # The template has no 'Table Grid' style: draw single-line borders directly instead
            logger.info("Style 'Table Grid' not in template; adding table borders manually.")
            tbl_pr = table._tbl.tblPr
            borders = OxmlElement('w:tblBorders')
            for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
                el = OxmlElement(f'w:{edge}')
                el.set(qn('w:val'), 'single')
                el.set(qn('w:sz'), '6')
                el.set(qn('w:space'), '0')
                el.set(qn('w:color'), '808080')
                borders.append(el)
            tbl_pr.append(borders)
        for r_idx, row_data in enumerate(table_data):
            row = table.add_row()
            for c_idx, cell_data in enumerate(row_data):
                cell = row.cells[c_idx]
                cell.text = str(cell_data)
                if r_idx == 0 and cell.paragraphs[0].runs:
                    cell.paragraphs[0].runs[0].bold = True
        return table

    def generate_report_from_template(self, template_path, output_path, data_placeholders, table_data=None):
        """
        Generates a Word document from a template, filling placeholders and optional table data.
        `data_placeholders` is a dictionary where keys are placeholders (e.g., '{{REPORT_TITLE}}') 
        and values are the replacement strings.
        `table_data` is a list of lists, where the first sub-list is the header row.
        """
        if not os.path.exists(template_path):
            logger.error(f"Word template not found: {template_path}")
            return False
        
        try:
            document = docx.Document(template_path)
            # Config keys may be written as "REPORT_TITLE" or "{{REPORT_TITLE}}"; the template uses {{...}}
            data_placeholders = {
                (k if k.startswith("{{") else "{{" + k + "}}"): v for k, v in data_placeholders.items()
            }
            
            # Replace placeholders in paragraphs
            for paragraph in document.paragraphs:
                for placeholder, value in data_placeholders.items():
                    self._replace_placeholder_in_paragraph(paragraph, placeholder, value)
            
            # Replace placeholders in table cells
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            for placeholder, value in data_placeholders.items():
                                self._replace_placeholder_in_paragraph(paragraph, placeholder, value)
            
            # Add dynamic table if provided
            if table_data:
                logger.info("Adding dynamic table to the document.")
                rows = len(table_data)
                cols = len(table_data[0]) if table_data else 0
                
                # Find a specific placeholder to insert table, or append to end
                table_anchor_found = False
                for paragraph in document.paragraphs:
                    if "{{TABLE_PLACEHOLDER}}" in paragraph.text:
                        table = self._build_table(document, table_data, cols)
                        # Move the table to where the placeholder paragraph is, then remove the placeholder
                        paragraph._p.addnext(table._tbl)
                        paragraph._p.getparent().remove(paragraph._p)
                        table_anchor_found = True
                        break  # Only replace first placeholder

                if not table_anchor_found:
                    logger.warning("Table placeholder '{{TABLE_PLACEHOLDER}}' not found in template. Appending table to end.")
                    self._build_table(document, table_data, cols)

            # Save the new document
            os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
            document.save(output_path)
            logger.info(f"Report generated from template and saved to '{output_path}'.")
            return True
            
        except Exception as e:
            logger.error(f"Error generating report from template: {e}")
            return False

    def extract_text_from_word(self, docx_path):
        """Extracts all text from a Word document."""
        if not os.path.exists(docx_path):
            logger.error(f"Word document not found: {docx_path}")
            return None
        
        try:
            document = docx.Document(docx_path)
            full_text = []
            
            for para in document.paragraphs:
                full_text.append(para.text)
            
            # Also extract text from tables
            for table in document.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        row_text.append(cell.text)
                    full_text.append("\t".join(row_text))  # Tab-separated cell values
            
            logger.info(f"Text extracted from Word document: {docx_path}")
            return "\n".join(full_text)
            
        except Exception as e:
            logger.error(f"Error extracting text from Word document '{docx_path}': {e}")
            return None